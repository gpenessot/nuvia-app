"""Livrables de l'étude : classeur Excel, note Word, synthèse PowerPoint.

`produire_livrables()` enchaîne tout.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd

try:  # utilisable comme module du paquet `outils` ou depuis le dossier lui-même
    from . import figures as F
    from .chargement import charger_tout
except ImportError:
    import figures as F
    from chargement import charger_tout


def _fr(valeur: float, decimales: int = 2) -> str:
    """Nombre à la française : virgule décimale. `_fr(0.3, 1)` donne « 0,3 »."""
    return f"{valeur:.{decimales}f}".replace(".", ",")


def _pct(valeur: float) -> str:
    """Pourcentage à la française : « 84 % », avec une espace insécable pour que
    Word ne renvoie jamais le signe seul en début de ligne."""
    return f"{valeur * 100:.0f} %"


def _pct_signe(valeur: float) -> str:
    """Écart en pourcentage, avec son signe : « +6 % », « -12 % »."""
    return f"{valeur * 100:+.0f} %"


# --------------------------------------------------------------------------
# Tables communes aux trois livrables
# --------------------------------------------------------------------------

def synthese_zonage(comparaison: pd.DataFrame, criteres: dict) -> pd.DataFrame:
    """Nombre de points par zone (en colonnes) et par phase (en lignes)."""
    ordre = [z["zone"] for z in criteres["zonage"]]
    table = pd.crosstab(comparaison["phase"], comparaison["zone"])
    return table.reindex(columns=[z for z in ordre if z in table.columns], fill_value=0)


def chiffres_cles(tout: dict, phase: int) -> dict:
    comparaison = tout["comparaison"].query("phase == @phase")
    ecart = tout["comparaison"]["ecart_relatif"]
    return {
        "listings": tout["calculs"]["fichier"].nunique(),
        "enregistrements": len(tout["calculs"]),
        "ecartes": int((~tout["calculs"]["exploitable"]).sum()),
        "mesures": len(tout["mesures"]),
        "sous_ld": int(tout["mesures"]["sous_limite_detection"].sum()),
        "ecart_median": ecart.median(),
        "ecart_dans_30pct": (ecart.abs() <= 0.30).mean(),
        "debit_max": comparaison["debit_retenu"].max(),
        "point_max": comparaison.loc[comparaison["debit_retenu"].idxmax(), "point"],
        "postes_non_conformes": int((~tout["postes"].query("phase == @phase")["conforme"]).sum()),
    }


def _table_postes(postes: pd.DataFrame, phase: int) -> pd.DataFrame:
    return (postes.query("phase == @phase")
            [["poste", "libelle", "point", "debit_retenu", "heures",
              "dose_individuelle_msv", "temps_max_h", "conforme"]]
            .rename(columns={"poste": "Poste", "libelle": "Libellé", "point": "Point",
                             "debit_retenu": "Débit (µSv/h)", "heures": "Durée (h)",
                             "dose_individuelle_msv": "Dose ind. (mSv)",
                             "temps_max_h": "Temps max (h)", "conforme": "Conforme"}))


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------

def ecrire_xlsx(tout: dict, chemin: str | Path) -> Path:
    """Classeur complet de l'étude, mis en forme.

    La couleur des zones est un format conditionnel Excel, pas une couleur posée
    cellule par cellule : si quelqu'un corrige une zone dans le fichier, la
    couleur suit.
    """
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    chemin = Path(chemin)
    calculs = tout["calculs"].drop(columns=["tally", "nps"])
    ecartes = calculs.loc[~calculs["exploitable"]].assign(
        raison=lambda d: d["tests_statistiques_ok"].map(
            {False: "tests statistiques MCNP non passés",
             True: "erreur relative > seuil"}))
    feuilles = {
        "Calculs": calculs,
        "Ecartes": ecartes,
        "Mesures": tout["mesures"],
        "Comparaison": tout["comparaison"][[
            "phase", "point", "local", "description", "x_cm", "y_cm",
            "debit_calcule", "erreur_relative", "debit_mesure", "sous_limite_detection",
            "ecart_relatif", "debit_retenu", "zone", "exploitable"]],
        "Postes": tout["postes"],
    }
    formats = {"debit": "0.00", "ecart": "0%", "erreur": "0.0%", "dose": "0.000",
               "heures": "0.0", "temps": "0.0"}
    entete = PatternFill("solid", start_color="1F3A5F")

    with pd.ExcelWriter(chemin, engine="openpyxl") as writer:
        for nom, table in feuilles.items():
            table.to_excel(writer, sheet_name=nom, index=False)
            ws = writer.sheets[nom]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            for i, colonne in enumerate(table.columns, start=1):
                lettre = get_column_letter(i)
                cellule = ws[f"{lettre}1"]
                cellule.font = Font(bold=True, color="FFFFFF")
                cellule.fill = entete
                cellule.alignment = Alignment(wrap_text=True, vertical="center")

                largeur = max(len(str(colonne)), table[colonne].astype(str).str.len().max())
                ws.column_dimensions[lettre].width = min(max(largeur + 2, 9), 40)

                prefixe = next((f for f in formats if colonne.startswith(f)), None)
                if prefixe:
                    for (cell,) in ws[f"{lettre}2:{lettre}{len(table) + 1}"]:
                        cell.number_format = formats[prefixe]

            if "zone" in table.columns:
                lettre = get_column_letter(list(table.columns).index("zone") + 1)
                plage = f"{lettre}2:{lettre}{len(table) + 1}"
                for zone in tout["criteres"]["zonage"]:
                    couleur = zone["couleur"].lstrip("#")
                    ws.conditional_formatting.add(plage, FormulaRule(
                        formula=[f'${lettre}2="{zone["zone"]}"'],
                        fill=PatternFill("solid", start_color=couleur, end_color=couleur)))

            if "conforme" in table.columns:
                lettre = get_column_letter(list(table.columns).index("conforme") + 1)
                ws.conditional_formatting.add(
                    f"{lettre}2:{lettre}{len(table) + 1}",
                    FormulaRule(formula=[f"${lettre}2=FALSE"],
                                font=Font(bold=True, color="C00000")))
    return chemin


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------

def _ombrer(cellule, couleur_hex: str) -> None:
    """Colore le fond d'une cellule Word. python-docx n'expose pas cette propriété :
    on écrit directement l'élément XML `w:shd` qu'utilise Word."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), couleur_hex.lstrip("#"))
    cellule._tc.get_or_add_tcPr().append(shd)


def _tableau_docx(doc, table: pd.DataFrame, couleurs: dict[str, str] | None = None,
                  formats: dict[str, str] | None = None):
    """Insère un DataFrame comme tableau Word. `couleurs` ombre les cellules dont
    la valeur est une clé du dictionnaire (typiquement le nom d'une zone)."""
    formats = formats or {}
    tableau = doc.add_table(rows=1, cols=len(table.columns))
    try:
        tableau.style = "Light Grid Accent 1"
    except KeyError:          # le modèle d'entreprise peut ne pas avoir ce style
        tableau.style = "Table Grid"

    for cellule, nom in zip(tableau.rows[0].cells, table.columns):
        cellule.text = str(nom)

    for _, ligne in table.iterrows():
        cellules = tableau.add_row().cells
        for cellule, (nom, valeur) in zip(cellules, ligne.items()):
            if isinstance(valeur, bool):
                texte = "oui" if valeur else "NON"
            elif isinstance(valeur, float):
                texte = format(valeur, formats.get(nom, ".2f")).replace(".", ",")
            else:
                texte = str(valeur)
            cellule.text = texte
            if couleurs and texte in couleurs:
                _ombrer(cellule, couleurs[texte])
    return tableau


def ecrire_docx(tout: dict, phase: int, images: dict[str, Path | None],
                chemin: str | Path, modele: str | Path | None = None) -> Path:
    """Note d'étude. `modele` : un .docx dont on reprend styles, en-tête et pied."""
    from docx import Document
    from docx.shared import Cm

    doc = Document(str(modele)) if modele else Document()
    criteres, comparaison = tout["criteres"], tout["comparaison"]
    c = chiffres_cles(tout, phase)
    libelle_phase = F.LIBELLES_PHASES[phase]

    doc.core_properties.title = f"Cartographie radiologique casemate C2 : phase {phase}"
    doc.core_properties.author = "Formation Python NUVIA 2026"

    doc.add_heading(f"Cartographie radiologique de la casemate C2", level=0)
    doc.add_paragraph(f"{libelle_phase} : note générée le {date.today():%d/%m/%Y}")

    def figure(nom: str, legende: str) -> None:
        if images.get(nom):
            doc.add_picture(str(images[nom]), width=Cm(16))
            doc.add_paragraph(legende, style="Caption")
        else:
            doc.add_paragraph(f"[Figure indisponible : {legende}]")

    doc.add_heading("1. Objet et données", level=1)
    doc.add_paragraph(
        f"La présente note établit la cartographie de débit de dose de la casemate C2 "
        f"et le zonage radiologique qui en découle, à l'issue de la {libelle_phase.lower()}. "
        f"Elle s'appuie sur {c['listings']} calculs MCNP ({c['enregistrements']} résultats "
        f"de détecteurs ponctuels) et sur {c['mesures']} mesures au radiamètre réparties en "
        f"trois campagnes.")
    doc.add_paragraph(
        f"{c['ecartes']} résultats de calcul ont été écartés, pour tests statistiques MCNP "
        f"non passés ou erreur relative supérieure à "
        f"{_pct(criteres['erreur_relative_max_acceptee'])}. {c['sous_ld']} mesures sont "
        f"inférieures à la limite de détection de l'appareil.")

    doc.add_heading("2. Qualification du modèle de calcul", level=1)
    doc.add_paragraph(
        f"L'écart relatif médian du calcul à la mesure est de {_pct_signe(c['ecart_median'])}, "
        f"et {_pct(c['ecart_dans_30pct'])} des points présentent un écart inférieur à 30 %. "
        f"Le modèle est jugé représentatif et peut être utilisé là où la mesure n'est "
        f"pas disponible.")
    figure("calcul_mesure", "Figure 1 : Débit calculé contre débit mesuré")

    doc.add_heading("3. Cartographie", level=1)
    doc.add_paragraph(
        f"Le débit de dose retenu en chaque point est l'enveloppe du calcul nominal et de "
        f"la mesure. Le maximum atteint {_fr(c['debit_max'], 1)} µSv/h au point "
        f"{c['point_max']}.")
    figure("plan", f"Figure 2 : Cartographie interpolée, {libelle_phase.lower()}")

    doc.add_heading("4. Zonage radiologique", level=1)
    zonage = synthese_zonage(comparaison, criteres).reset_index()
    zonage["phase"] = zonage["phase"].map(F.LIBELLES_PHASES)
    _tableau_docx(doc, zonage.rename(columns={"phase": "Phase"}))
    doc.add_paragraph()
    figure("zonage", "Figure 3 : Nombre de points par zone et par phase")

    doc.add_heading("5. Postes de travail", level=1)
    doc.add_paragraph(
        f"Dose individuelle prévisionnelle par poste, comparée à la contrainte de chantier "
        f"de {_fr(criteres['contrainte_chantier_mSv'], 1)} mSv. Chaque intervenant est supposé "
        f"présent pendant toute la durée du poste, au débit du point de référence.")
    _tableau_docx(doc, _table_postes(tout["postes"], phase), couleurs={"NON": "F4B6B6"},
                  formats={"Dose ind. (mSv)": ".3f", "Durée (h)": ".1f",
                           "Temps max (h)": ".1f"})
    doc.add_paragraph()
    if c["postes_non_conformes"]:
        doc.add_paragraph(
            f"{c['postes_non_conformes']} poste(s) dépasse(nt) la contrainte : une "
            f"optimisation est nécessaire avant intervention (écran, réduction du temps "
            f"de présence, éloignement).")
    else:
        doc.add_paragraph("Tous les postes respectent la contrainte de chantier.")

    doc.add_heading("6. Hypothèses et limites", level=1)
    for texte in [
        "Le débit retenu est l'enveloppe max(calcul nominal, mesure) : choix conservatif.",
        "Une mesure inférieure à la limite de détection est remplacée par le calcul.",
        "La cartographie est interpolée par pondération inverse de la distance sur le "
        "logarithme du débit, local par local : elle n'a pas de valeur entre deux points "
        "éloignés.",
        f"Seuils de zonage : {criteres['reference']}. {criteres['avertissement']}",
    ]:
        doc.add_paragraph(texte, style="List Bullet")

    chemin = Path(chemin)
    doc.save(str(chemin))
    return chemin


# --------------------------------------------------------------------------
# PowerPoint
# --------------------------------------------------------------------------

def ecrire_pptx(tout: dict, phase: int, images: dict[str, Path | None],
                chemin: str | Path) -> Path:
    """Synthèse en cinq diapositives, format 16:9."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Cm, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    titre_seul = prs.slide_layouts[5]
    c = chiffres_cles(tout, phase)
    criteres = tout["criteres"]

    def diapo(titre: str):
        slide = prs.slides.add_slide(titre_seul)
        slide.shapes.title.text = titre
        slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
        return slide

    def image(slide, nom: str) -> None:
        if images.get(nom):
            photo = slide.shapes.add_picture(str(images[nom]), 0, Cm(4), height=Cm(14.5))
            photo.left = int((prs.slide_width - photo.width) / 2)
        else:
            boite = slide.shapes.add_textbox(Cm(2.5), Cm(8), Cm(28), Cm(3))
            boite.text_frame.text = "Figure indisponible (export PNG impossible sur ce poste)"

    # 1 : titre
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Casemate C2 : cartographie radiologique"
    slide.placeholders[1].text = f"{F.LIBELLES_PHASES[phase]}\n{date.today():%d/%m/%Y}"

    # 2 : chiffres clés
    slide = diapo("Ce qu'il faut retenir")
    cadre = slide.shapes.add_textbox(Cm(2.5), Cm(4.5), Cm(29), Cm(13)).text_frame
    cadre.word_wrap = True
    lignes = [
        f"Débit maximal : {_fr(c['debit_max'], 1)} µSv/h (point {c['point_max']})",
        f"Modèle qualifié : écart médian {_pct_signe(c['ecart_median'])}, "
        f"{_pct(c['ecart_dans_30pct'])} des points à ±30 %",
        f"{c['ecartes']} calculs écartés sur {c['enregistrements']} (convergence)",
        f"Postes hors contrainte ({_fr(criteres['contrainte_chantier_mSv'], 1)} mSv) : "
        f"{c['postes_non_conformes']}",
    ]
    for i, texte in enumerate(lignes):
        paragraphe = cadre.paragraphs[0] if i == 0 else cadre.add_paragraph()
        paragraphe.text = f"•  {texte}"
        paragraphe.font.size = Pt(24)
        paragraphe.space_after = Pt(18)

    # 3 (carte, 4) zonage
    image(diapo(f"Cartographie : {F.LIBELLES_PHASES[phase]}"), "plan")
    image(diapo("Zonage réglementaire par phase"), "zonage")

    # 5 : postes
    slide = diapo("Postes de travail")
    table = _table_postes(tout["postes"], phase).drop(columns=["Temps max (h)"])
    forme = slide.shapes.add_table(len(table) + 1, len(table.columns),
                                   Cm(1.5), Cm(4.5), Cm(31), Cm(1.2) * (len(table) + 1))
    grille = forme.table
    for j, nom in enumerate(table.columns):
        grille.cell(0, j).text = nom
    for i, (_, ligne) in enumerate(table.iterrows(), start=1):
        for j, valeur in enumerate(ligne):
            cellule = grille.cell(i, j)
            if isinstance(valeur, bool):
                cellule.text = "oui" if valeur else "NON"
                if not valeur:
                    cellule.fill.solid()
                    cellule.fill.fore_color.rgb = RGBColor.from_string("F4B6B6")
            else:
                cellule.text = (f"{valeur:.3g}".replace(".", ",") if isinstance(valeur, float)
                                else str(valeur))
    for ligne in grille.rows:
        for cellule in ligne.cells:
            cellule.text_frame.paragraphs[0].font.size = Pt(14)

    chemin = Path(chemin)
    prs.save(str(chemin))
    return chemin


# --------------------------------------------------------------------------
# Tout d'un coup
# --------------------------------------------------------------------------

def produire_livrables(racine: str | Path | None = None, phase: int = 2,
                       sortie: str | Path = "livrables",
                       modele: str | Path | None = None) -> dict[str, Path]:
    """Charge les données, trace, exporte, écrit les trois livrables.

    Renvoie le chemin de chaque fichier produit. Une figure dont l'export PNG
    échoue est remplacée par une mention dans les documents, sans interrompre
    la production.
    """
    sortie = Path(sortie)
    sortie.mkdir(parents=True, exist_ok=True)
    tout = charger_tout(racine)
    comparaison, criteres, plan = tout["comparaison"], tout["criteres"], tout["plan"]

    figures = {
        "calcul_mesure": (F.fig_calcul_mesure(comparaison), 800, 700),
        "plan": (F.fig_plan(plan, comparaison, phase, criteres), 1100, 620),
        "zonage": (F.fig_zonage(comparaison, criteres), 900, 480),
    }
    carte_html = sortie / f"carte_phase{phase}.html"
    figures["plan"][0].write_html(carte_html, include_plotlyjs=True)

    images: dict[str, Path | None] = {}
    with F.serveur_export():
        for nom, (fig, largeur, hauteur) in figures.items():
            # Dans un document, la légende Word ou le titre de diapositive fait foi :
            # le titre Plotly ferait doublon.
            fig.update_layout(title=None, margin=dict(t=30))
            chemin = sortie / f"figure_{nom}.png"
            images[nom] = chemin if F.exporter_png(fig, chemin, largeur, hauteur) else None

    return {
        "xlsx": ecrire_xlsx(tout, sortie / "cartographie.xlsx"),
        "docx": ecrire_docx(tout, phase, images, sortie / f"note_etude_phase{phase}.docx",
                            modele),
        "pptx": ecrire_pptx(tout, phase, images, sortie / f"synthese_phase{phase}.pptx"),
        "html": carte_html,
    }


def _autotest() -> None:
    import tempfile
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation

    with tempfile.TemporaryDirectory() as dossier:
        modele = Path(__file__).resolve().parent.parent / "data" / "modele_note.docx"
        fichiers = produire_livrables(phase=0, sortie=dossier,
                                      modele=modele if modele.exists() else None)
        assert all(p.exists() and p.stat().st_size > 0 for p in fichiers.values()), fichiers

        classeur = load_workbook(fichiers["xlsx"])
        assert classeur.sheetnames == ["Calculs", "Ecartes", "Mesures", "Comparaison",
                                       "Postes"], classeur.sheetnames
        assert len(classeur["Comparaison"].conditional_formatting) >= 1

        note = Document(str(fichiers["docx"]))
        titres = [p.text for p in note.paragraphs if p.style.name.startswith("Heading")]
        assert len(titres) == 6, titres
        assert len(note.inline_shapes) == 3, "les trois figures doivent être insérées"
        # En phase 0, un poste dépasse la contrainte : la note doit le dire.
        assert any("dépasse" in p.text for p in note.paragraphs)

        diapos = Presentation(str(fichiers["pptx"]))
        assert len(diapos.slides) == 5

    print(f"autotest : ok : xlsx 5 feuilles, docx 6 sections et 3 figures, pptx 5 diapositives")


if __name__ == "__main__":
    import sys
    sys.path.insert(0, str(Path(__file__).parent))
    _autotest()
